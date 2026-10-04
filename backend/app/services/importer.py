# backend/app/services/importer.py
"""Importação de leituras a partir de Excel (.xlsx) ou CSV.

- Reconhece cabeçalhos em português ("Temp. óleo (°C)") ou os nomes técnicos
  ("oil_temperature"); maiúsculas, acentos e unidades entre parênteses são ignorados.
- Valida cada linha com as mesmas regras do registo manual (OperationalReadingCreate).
- Rejeita leituras duplicadas (mesmo grupo e data/hora) no ficheiro ou na base de dados.
"""
import csv
import io
import json
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Any

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import select, tuple_
from sqlalchemy.orm import Session

from app.models.group import Group
from app.models.operational_reading import OperationalReading
from app.schemas.reading import OperationalReadingCreate

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20_000
# Tolerância para relógios dessincronizados
FUTURE_TOLERANCE = timedelta(hours=1)

# Cabeçalhos do modelo (ordem e texto apresentados ao utilizador)
TEMPLATE_COLUMNS: list[tuple[str, str]] = [
    ("group_code", "Grupo"),
    ("timestamp", "Data e hora"),
    ("operator", "Operador"),
    ("status", "Estado"),
    ("power", "Potência ativa (MW)"),
    ("reactive_power", "Potência reativa (MVAr)"),
    ("frequency", "Frequência (Hz)"),
    ("voltage", "Tensão (V)"),
    ("current_r", "Corrente R (A)"),
    ("current_s", "Corrente S (A)"),
    ("current_t", "Corrente T (A)"),
    ("power_factor", "Fator de potência"),
    ("oil_temperature", "Temp. óleo (°C)"),
    ("oil_pressure", "Pressão óleo (bar)"),
    ("water_temperature", "Temp. água (°C)"),
    ("intake_temperature", "Temp. admissão (°C)"),
    ("exhaust_temp_left", "Temp. escape esq. (°C)"),
    ("exhaust_temp_right", "Temp. escape dir. (°C)"),
    ("operating_hours", "Horas de funcionamento (h)"),
    ("energy_produced", "Energia acumulada (MWh)"),
]
FIELD_LABELS = dict(TEMPLATE_COLUMNS)
REQUIRED_FIELDS = ("group_code", "timestamp", "status")
NUMERIC_FIELDS = [f for f, _ in TEMPLATE_COLUMNS if f not in ("group_code", "timestamp", "operator", "status")]

# Sinónimos aceites (já normalizados)
ALIASES: dict[str, list[str]] = {
    "group_code": ["grupo", "group", "codigo_grupo", "codigo", "gg"],
    "timestamp": ["data_e_hora", "data_hora", "datahora", "data", "date", "datetime"],
    "operator": ["operador", "responsavel"],
    "status": ["estado", "estado_operacional", "situacao"],
    "power": ["potencia_ativa", "potencia", "potencia_activa", "p"],
    "reactive_power": ["potencia_reativa", "potencia_reactiva", "q"],
    "frequency": ["frequencia", "freq"],
    "voltage": ["tensao", "v"],
    "current_r": ["corrente_r", "ir", "i_r"],
    "current_s": ["corrente_s", "is", "i_s"],
    "current_t": ["corrente_t", "it", "i_t"],
    "power_factor": ["fator_de_potencia", "fator_potencia", "factor_de_potencia", "fp", "cos_phi", "cosphi"],
    "oil_temperature": ["temp_oleo", "temperatura_oleo", "temperatura_do_oleo", "temp_do_oleo"],
    "oil_pressure": ["pressao_oleo", "pressao_do_oleo"],
    "water_temperature": ["temp_agua", "temperatura_agua", "temperatura_da_agua", "temp_da_agua"],
    "intake_temperature": ["temp_admissao", "temperatura_admissao", "temperatura_de_admissao"],
    "exhaust_temp_left": ["temp_escape_esq", "temp_escape_esquerdo", "temperatura_escape_esq", "escape_esq"],
    "exhaust_temp_right": ["temp_escape_dir", "temp_escape_direito", "temperatura_escape_dir", "escape_dir"],
    "operating_hours": ["horas_de_funcionamento", "horas_funcionamento", "horas", "horimetro"],
    "energy_produced": ["energia_acumulada", "energia_produzida", "energia", "contador_energia"],
}

STATUS_VALUES = {
    "em operacao": "Em operação",
    "operacao": "Em operação",
    "em operação": "Em operação",
    "operando": "Em operação",
    "operacional": "Em operação",
    "parado": "Parado",
    "manutencao": "Manutenção",
    "em manutencao": "Manutenção",
}

DATE_FORMATS = (
    "%d/%m/%Y %H:%M",
    "%d/%m/%Y %H:%M:%S",
    "%d-%m-%Y %H:%M",
    "%d-%m-%Y %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%dT%H:%M",
    "%Y-%m-%dT%H:%M:%S",
)


def normalize(text: Any) -> str:
    """'Temp. óleo (°C)' -> 'temp_oleo'"""
    s = str(text or "")
    s = re.sub(r"\(.*?\)", "", s)  # unidades
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    s = re.sub(r"[^a-z0-9]+", "_", s.lower())
    return s.strip("_")


HEADER_MAP: dict[str, str] = {}
for _field, _label in TEMPLATE_COLUMNS:
    HEADER_MAP[_field] = _field
    HEADER_MAP[normalize(_label)] = _field
    for _alias in ALIASES.get(_field, []):
        HEADER_MAP[_alias] = _field


class ImportFileError(ValueError):
    """Erro que impede a leitura do ficheiro inteiro."""


@dataclass
class RowError:
    row_number: int
    column: str | None
    message: str
    raw: dict[str, Any]


@dataclass
class ParsedFile:
    file_type: str
    headers: list[str]
    mapping: dict[str, str | None]  # cabeçalho original -> campo
    rows: list[tuple[int, dict[str, Any]]]  # (nº da linha no ficheiro, valores por cabeçalho)


@dataclass
class ValidationResult:
    valid: list[tuple[int, OperationalReadingCreate]] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)

    @property
    def error_rows(self) -> int:
        return len({e.row_number for e in self.errors})


# ------------------------------------------------------------------ leitura
def read_file(filename: str, content: bytes) -> ParsedFile:
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError("O ficheiro excede 5 MB.")
    name = filename.lower()
    if name.endswith(".xlsx"):
        file_type, table = "EXCEL", _read_xlsx(content)
    elif name.endswith(".csv"):
        file_type, table = "CSV", _read_csv(content)
    else:
        raise ImportFileError("Formato não suportado. Use .xlsx ou .csv.")

    # Primeira linha não vazia = cabeçalhos
    table = [(i, r) for i, r in table if any(str(c).strip() for c in r if c is not None)]
    if not table:
        raise ImportFileError("O ficheiro está vazio.")
    _, header_row = table[0]
    headers = [str(h).strip() if h is not None else "" for h in header_row]

    data_rows = table[1:]
    if len(data_rows) > MAX_ROWS:
        raise ImportFileError(f"O ficheiro tem mais de {MAX_ROWS} linhas.")

    mapping = {h: HEADER_MAP.get(normalize(h)) for h in headers if h}
    rows = [
        (line, {headers[c]: v for c, v in enumerate(values) if c < len(headers) and headers[c]})
        for line, values in data_rows
    ]
    return ParsedFile(file_type=file_type, headers=[h for h in headers if h], mapping=mapping, rows=rows)


def _read_xlsx(content: bytes) -> list[tuple[int, list[Any]]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # ficheiro corrompido / não é xlsx
        raise ImportFileError("Não foi possível abrir o ficheiro Excel.") from exc
    ws = wb.worksheets[0]
    return [(i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), start=1)]


def _read_csv(content: bytes) -> list[tuple[int, list[Any]]]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        class dialect(csv.excel):  # noqa: N801
            delimiter = ";" if text.count(";") > text.count(",") else ","
    return [(i, row) for i, row in enumerate(csv.reader(io.StringIO(text), dialect), start=1)]


# ------------------------------------------------------------------ conversões
def parse_datetime(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None, microsecond=0)
    text = str(value or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError("Data/hora inválida (use dd/mm/aaaa hh:mm)")


def parse_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if not text or text in ("-", "—"):
        return None
    # "1.234,5" -> 1234.5 ; "1234,5" -> 1234.5
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        raise ValueError("Número inválido") from None


def parse_status(value: Any) -> str:
    key = unicodedata.normalize("NFKD", str(value or "").strip().lower()).encode("ascii", "ignore").decode()
    if key not in STATUS_VALUES:
        raise ValueError("Estado inválido (use: Em operação, Parado ou Manutenção)")
    return STATUS_VALUES[key]


def _jsonable(raw: dict[str, Any]) -> dict[str, Any]:
    return {k: (v.isoformat(sep=" ") if isinstance(v, datetime) else v) for k, v in raw.items() if v not in (None, "")}


# ------------------------------------------------------------------ validação
def missing_required(parsed: ParsedFile) -> list[str]:
    present = set(parsed.mapping.values())
    return [FIELD_LABELS[f] for f in REQUIRED_FIELDS if f not in present]


def validate(parsed: ParsedFile, db: Session, default_operator: str) -> ValidationResult:
    result = ValidationResult()
    groups = {g.code.upper(): g for g in db.scalars(select(Group))}
    seen: set[tuple[int, datetime]] = set()
    candidates: list[tuple[int, OperationalReadingCreate, int, dict]] = []

    field_to_header = {f: h for h, f in parsed.mapping.items() if f}

    for line, raw in parsed.rows:
        values = {f: raw.get(h) for f, h in field_to_header.items()}
        errors: list[tuple[str | None, str]] = []

        code = str(values.get("group_code") or "").strip().upper()
        group = groups.get(code)
        if not code:
            errors.append(("group_code", "Grupo em falta"))
        elif group is None:
            errors.append(("group_code", f"Grupo não encontrado: {code}"))

        data: dict[str, Any] = {"group_code": code}
        try:
            data["timestamp"] = parse_datetime(values.get("timestamp"))
            if data["timestamp"] > datetime.now() + FUTURE_TOLERANCE:
                errors.append(("timestamp", "Data/hora no futuro"))
        except ValueError as exc:
            errors.append(("timestamp", str(exc)))
        try:
            data["status"] = parse_status(values.get("status"))
        except ValueError as exc:
            errors.append(("status", str(exc)))

        operator = str(values.get("operator") or "").strip()
        data["operator"] = operator or default_operator

        for f in NUMERIC_FIELDS:
            if f not in values:
                continue
            try:
                data[f] = parse_number(values[f])
            except ValueError as exc:
                errors.append((f, str(exc)))

        reading = None
        if not errors:
            try:
                reading = OperationalReadingCreate(**data)
            except ValidationError as exc:
                for e in exc.errors():
                    loc = e["loc"][0] if e["loc"] else None
                    errors.append((str(loc) if loc else None, _translate(e)))

        if reading is not None and group is not None:
            key = (group.id, reading.timestamp)
            if key in seen:
                errors.append(("timestamp", "Leitura duplicada no ficheiro (mesmo grupo e data/hora)"))
            seen.add(key)

        if errors:
            for col, msg in errors:
                result.errors.append(RowError(line, FIELD_LABELS.get(col, col) if col else None, msg, _jsonable(raw)))
        else:
            candidates.append((line, reading, group.id, raw))

    # Duplicados já existentes na base de dados
    keys = [(gid, r.timestamp) for _, r, gid, _ in candidates]
    existing: set[tuple[int, datetime]] = set()
    for i in range(0, len(keys), 1000):
        chunk = keys[i : i + 1000]
        existing.update(
            (gid, ts)
            for gid, ts in db.execute(
                select(OperationalReading.group_id, OperationalReading.timestamp).where(
                    tuple_(OperationalReading.group_id, OperationalReading.timestamp).in_(chunk)
                )
            ).all()
        )

    for line, reading, gid, raw in candidates:
        if (gid, reading.timestamp) in existing:
            result.errors.append(
                RowError(line, FIELD_LABELS["timestamp"], "Já existe uma leitura deste grupo nesta data/hora", _jsonable(raw))
            )
        else:
            result.valid.append((line, reading))

    result.errors.sort(key=lambda e: e.row_number)
    return result


def _translate(error: dict) -> str:
    kind = error.get("type", "")
    ctx = error.get("ctx") or {}
    if kind == "greater_than_equal":
        return f"Valor abaixo do mínimo permitido ({ctx.get('ge')})"
    if kind == "less_than_equal":
        return f"Valor acima do máximo permitido ({ctx.get('le')})"
    if kind == "missing":
        return "Campo obrigatório em falta"
    return error.get("msg", "Valor inválido")


def raw_json(raw: dict[str, Any]) -> str:
    return json.dumps(raw, ensure_ascii=False, default=str)


# ------------------------------------------------------------------ modelo
def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leituras"
    ws.append([label for _, label in TEMPLATE_COLUMNS])
    ws.append(["GG1", "03/10/2026 08:00", "Ana Júlia", "Em operação", 0.82, 0.25, 50.02, 401, 1180, 1175, 1182,
               0.95, 84.5, 4.8, 82.1, 31.0, 465, 470, 4120.5, 25310.2])
    ws.append(["GG4", "03/10/2026 08:00", "Ana Júlia", "Parado"] + [None] * (len(TEMPLATE_COLUMNS) - 4))
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 2)
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Instruções")
    for line in [
        "Importação de leituras operacionais",
        "",
        "ATENÇÃO: apague as linhas de exemplo (GG1 e GG4) antes de importar.",
        "",
        "Colunas obrigatórias: Grupo, Data e hora, Estado.",
        "Data e hora: dd/mm/aaaa hh:mm (ou uma célula de data do Excel).",
        "Estado: Em operação, Parado ou Manutenção.",
        "Números: aceita vírgula ou ponto decimal. Medições vazias são permitidas.",
        "Operador: se vazio, é usado o utilizador que faz a importação.",
        "Leituras repetidas (mesmo grupo e data/hora) são rejeitadas.",
        "Apenas a primeira folha do ficheiro é importada.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 90

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
